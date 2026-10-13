import io

import pandas as pd
import pytest
from app.core.interactive import basic_cleanup
from app.main import app, rate, store
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook


@pytest.fixture
def client(tmp_path):
    store.root = tmp_path
    store.jobs.clear()
    rate.clear()
    return TestClient(app)


def upload(client):
    response = client.post(
        "/api/files/upload",
        files={
            "file": (
                "editor.csv",
                b"Customer_ID,Name,Birth_Date,Salary,Email\n001, mary david ,2025-01-02,$100,invalid-email\n002,MARY DAVID,05/10/2025,200,mary@example.com\n003,Other,88-99-2089,NA,other@example.com\n",
                "text/csv",
            )
        },
    )
    assert response.status_code == 200, response.text
    return response.json()["file_id"]


def test_basic_cleanup_preserves_identifiers_acronyms_and_ambiguous_dates():
    df = pd.DataFrame(
        {
            "Customer_ID": ["001", "002", "003"],
            "Name": [" mary david ", "MARY DAVID", "Other"],
            "Department": ["IT", "it", "HR"],
            "Birth_Date": ["2025-01-02", "05/10/2025", "88-99-2089"],
            "Email": ["Mary@example.com", "MARY@example.com", "x@y.com"],
        }
    )
    original = df.copy(deep=True)
    cleaned, logs = basic_cleanup(df)
    assert cleaned["Name"].tolist() == ["Mary David", "Mary David", "Other"]
    assert cleaned["Department"].tolist() == ["IT", "IT", "HR"]
    assert cleaned["Customer_ID"].tolist() == ["001", "002", "003"]
    assert cleaned["Birth_Date"].tolist() == original["Birth_Date"].tolist()
    assert cleaned["Email"].tolist() == original["Email"].tolist()
    assert any(l["modified_values"] for l in logs)
    pd.testing.assert_frame_equal(df, original)


def test_editor_annotations_preview_confirm_export_and_access(client):
    fid = upload(client)
    started = client.post(f"/api/files/{fid}/editor/start").json()
    assert started["preview"]["rows"][0][1] == "Mary David"
    assert started["preview"]["cell_issues"][0][4][0]["kind"] == "email"
    assert any(
        i["kind"] == "ambiguous_date" for i in started["preview"]["cell_issues"][1][2]
    )
    assert any(
        i["kind"] == "invalid_date" for i in started["preview"]["cell_issues"][2][2]
    )
    assert started["preview"]["cell_issues"][2][3][0]["kind"] == "missing"
    body = {
        "revision": 0,
        "column": "Birth_Date",
        "row_index": 2,
        "action": "replace",
        "value": "2025-10-06",
    }
    preview = client.post(f"/api/files/{fid}/editor/fix", json=body)
    assert preview.status_code == 200, preview.text
    assert preview.json()["proposal"]["changed_cells"] == 1
    assert (
        client.get(f"/api/files/{fid}/editor").json()["preview"]["rows"][2][2]
        == "88-99-2089"
    )
    body["confirmed"] = True
    assert client.post(f"/api/files/{fid}/editor/fix", json=body).status_code == 200
    after = client.get(f"/api/files/{fid}/editor").json()
    assert after["revision"] == 1
    assert not after["preview"]["cell_issues"][2][2]
    assert (
        client.get(f"/api/files/{fid}/preview").json()["rows"][0][1] == " mary david "
    )
    assert client.post(f"/api/files/{fid}/editor/fix", json=body).status_code == 409
    outsider = TestClient(app)
    assert outsider.get(f"/api/files/{fid}/editor").status_code == 404
    assert outsider.get(f"/api/files/{fid}/editor/download").status_code == 404
    csv = client.get(f"/api/files/{fid}/editor/download?format=csv")
    assert "2025-10-06" in csv.text and "001" in csv.text
    book = load_workbook(
        io.BytesIO(client.get(f"/api/files/{fid}/editor/download?format=xlsx").content)
    )
    assert book.active["A2"].value == "001"
    assert (
        client.get(f"/api/files/{fid}/editor/download?format=report").json()[
            "confirmed_changes"
        ][0]["action"]
        == "replace"
    )
    assert client.post(f"/api/files/{fid}/editor/start").json()["revision"] == 1
    assert client.delete(f"/api/files/{fid}").status_code == 200
    assert client.get(f"/api/files/{fid}/editor").status_code == 404


def test_column_currency_and_cell_imputation_keep_other_cells_unchanged(client):
    fid = upload(client)
    client.post(f"/api/files/{fid}/editor/start")
    body = {
        "revision": 0,
        "column": "Salary",
        "row_index": 0,
        "action": "normalize_currency",
        "scope": "column",
        "parameters": {"action": "remove"},
        "confirmed": True,
    }
    response = client.post(f"/api/files/{fid}/editor/fix", json=body)
    assert response.status_code == 200, response.text
    body.update(
        revision=1,
        row_index=2,
        scope="cell",
        action="impute_missing_values",
        parameters={"method": "mean"},
    )
    assert client.post(f"/api/files/{fid}/editor/fix", json=body).status_code == 200
    values = client.get(f"/api/files/{fid}/editor").json()["preview"]["rows"]
    assert [r[3] for r in values] == ["100", "200", 150]
    assert values[0][4] == "invalid-email"


def test_invalid_editor_requests_and_formula_safe_export(client):
    fid = upload(client)
    assert client.get(f"/api/files/{fid}/editor").status_code == 409
    client.post(f"/api/files/{fid}/editor/start")
    body = {
        "revision": 0,
        "column": "Name",
        "row_index": 0,
        "action": "replace",
        "scope": "column",
        "value": "X",
        "confirmed": True,
    }
    assert client.post(f"/api/files/{fid}/editor/fix", json=body).status_code == 422
    body.update(scope="cell", row_index=999)
    assert client.post(f"/api/files/{fid}/editor/fix", json=body).status_code == 422
    body.update(row_index=0, value="=1+1")
    assert client.post(f"/api/files/{fid}/editor/fix", json=body).status_code == 200
    assert "'=1+1" in client.get(f"/api/files/{fid}/editor/download?format=csv").text
    assert client.get(f"/api/files/{fid}/editor?page=2&page_size=1").json()["preview"][
        "row_indices"
    ] == [1]


def test_editor_accepts_excel_and_retains_native_date_and_id(client):
    from datetime import datetime

    book = Workbook()
    book.active.append(["Customer_ID", "Name", "Birth_Date"])
    book.active.append(["001", " mary david ", datetime(2005, 3, 16)])
    buffer = io.BytesIO()
    book.save(buffer)
    uploaded = client.post(
        "/api/files/upload",
        files={
            "file": (
                "editor.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 200, uploaded.text
    fid = uploaded.json()["file_id"]
    response = client.post(f"/api/files/{fid}/editor/start")
    assert response.status_code == 200, response.text
    row = response.json()["preview"]["rows"][0]
    assert row[:2] == ["001", "Mary David"]
    assert row[2].startswith("2005-03-16")


def test_editor_column_percentage_and_confirmed_duplicate_removal(client):
    data = b"Name,Rate\nA,10%\nB,20\nB,20\n"
    fid = client.post(
        "/api/files/upload", files={"file": ("notation.csv", data, "text/csv")}
    ).json()["file_id"]
    client.post(f"/api/files/{fid}/editor/start")
    body = {
        "revision": 0,
        "column": "Rate",
        "row_index": 0,
        "scope": "cell",
        "action": "normalize_percentage",
        "parameters": {"action": "remove"},
        "confirmed": True,
    }
    assert client.post(f"/api/files/{fid}/editor/fix", json=body).status_code == 200
    rows = client.get(f"/api/files/{fid}/editor").json()["preview"]["rows"]
    assert [r[1] for r in rows] == ["10", "20", "20"]
    body.update(
        revision=1,
        column="Name",
        scope="column",
        action="remove_duplicates",
        parameters={},
        confirmed=False,
    )
    preview = client.post(f"/api/files/{fid}/editor/fix", json=body).json()["proposal"]
    assert preview["changed_cells"] == 1
    assert client.get(f"/api/files/{fid}/editor").json()["preview"]["total_rows"] == 3
    body["confirmed"] = True
    assert client.post(f"/api/files/{fid}/editor/fix", json=body).status_code == 200
    after = client.get(f"/api/files/{fid}/editor?page=9&page_size=1").json()["preview"]
    assert after["total_rows"] == 2 and after["page"] == 2
    assert after["row_indices"] == [1]
    assert not after["cell_issues"][0][0]


def test_editor_date_interpretation_and_new_whitespace_annotation(client):
    fid = upload(client)
    client.post(f"/api/files/{fid}/editor/start")
    body = {
        "revision": 0,
        "column": "Birth_Date",
        "row_index": 1,
        "scope": "cell",
        "action": "normalize_dates",
        "parameters": {"interpretation": "day_first"},
        "confirmed": True,
    }
    assert client.post(f"/api/files/{fid}/editor/fix", json=body).status_code == 200
    data = client.get(f"/api/files/{fid}/editor").json()
    assert data["preview"]["rows"][1][2] == "2025-10-05"
    assert data["preview"]["rows"][2][2] == "88-99-2089"
    body.update(
        revision=1,
        column="Name",
        row_index=0,
        action="replace",
        value=" Mary ",
        parameters={},
    )
    assert client.post(f"/api/files/{fid}/editor/fix", json=body).status_code == 200
    data = client.get(f"/api/files/{fid}/editor").json()
    assert any(i["kind"] == "whitespace" for i in data["preview"]["cell_issues"][0][1])
    body.update(revision=2, action="trim_whitespace")
    assert client.post(f"/api/files/{fid}/editor/fix", json=body).status_code == 200
    assert (
        client.get(f"/api/files/{fid}/editor").json()["preview"]["rows"][0][1] == "Mary"
    )
