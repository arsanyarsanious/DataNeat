import io
import time

import pytest
from app.main import app, rate, store
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook


@pytest.fixture
def client(tmp_path):
    store.root = tmp_path
    store.jobs.clear()
    rate.clear()
    # Avoid shutting down the shared executor between tests.
    return TestClient(app)


def upload(
    client,
    data=b"ID,City,Birth_Date,Notes\n001, Cairo ,05/10/2025,=1+1\n002,cairo,15/05/2025,normal\n",
):
    response = client.post(
        "/api/files/upload", files={"file": ("customers.csv", data, "text/csv")}
    )
    assert response.status_code == 200, response.text
    return response.json()


def wait_job(client, job):
    for _ in range(100):
        response = client.get(f"/api/jobs/{job}").json()
        if response["status"] == "completed":
            return
        assert response["status"] != "failed", response
        time.sleep(0.02)
    pytest.fail("Job did not complete")


def test_real_workflow_and_authorization(client):
    uploaded = upload(client)
    fid = uploaded["file_id"]
    assert client.get(f"/api/files/{fid}/profile").json()["rows"] == 2
    assert client.get(f"/api/files/{fid}/issues").status_code == 200
    assert client.get(f"/api/files/{fid}/cleaning-plan").status_code == 200
    outsider = TestClient(app)
    assert outsider.get(f"/api/files/{fid}/preview").status_code == 404
    plan = uploaded["plan"]
    for op in plan:
        op["enabled"] = True
    body = {"operations": plan}
    trial = client.post(f"/api/files/{fid}/preview-cleaning", json=body)
    assert trial.status_code == 200, trial.text
    assert trial.json()["after"]["rows"][0][2] == "05/10/2025"
    response = client.post(f"/api/files/{fid}/clean", json=body)
    assert response.status_code == 202
    job = response.json()["job_id"]
    wait_job(client, job)
    assert outsider.get(f"/api/jobs/{job}/download").status_code == 404
    results = client.get(f"/api/jobs/{job}/results").json()
    assert results["preview"]["rows"][0][0] == "001"
    assert results["preview"]["rows"][0][1] == "Cairo"
    assert results["validation"]["unresolved_issues"] > 0
    csv = client.get(f"/api/jobs/{job}/download?format=csv")
    assert "'=1+1" in csv.text and "001" in csv.text
    excel = client.get(f"/api/jobs/{job}/download?format=xlsx")
    workbook = load_workbook(io.BytesIO(excel.content))
    assert workbook.active["A2"].value == "001"
    assert workbook.active["D2"].data_type != "f"
    assert client.get(f"/api/jobs/{job}/download?format=report").json()["validation"]
    assert client.get(f"/api/files/{fid}/preview").json()["rows"][0][1] == " Cairo "
    assert client.delete(f"/api/files/{fid}").status_code == 200
    assert client.get(f"/api/jobs/{job}").status_code == 404
    assert client.get(f"/api/files/{fid}/profile").status_code == 404


def test_excel_sheets(client):
    workbook = Workbook()
    workbook.active.title = "First"
    workbook.active.append(["ID", "Date"])
    workbook.active.append(["001", "05/10/2025"])
    other = workbook.create_sheet("Second")
    other.append(["Name", "Name"])
    other.append(["Alice", "Bob"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = client.post(
        "/api/files/upload",
        files={"file": ("book.xlsx", buffer.getvalue())},
        data={"sheet": "Second"},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["metadata"]["sheets"] == ["First", "Second"]
    assert result["preview"]["columns"] == ["Name", "Name_2"]
    assert result["preview"]["rows"] == [["Alice", "Bob"]]


def test_validation_limits_and_origin(client):
    for name, data in [
        ("bad.txt", b"abc"),
        ("empty.csv", b""),
        ("invalid.xlsx", b"not a workbook"),
    ]:
        assert (
            client.post("/api/files/upload", files={"file": (name, data)}).status_code
            == 422
        )
    assert (
        client.post(
            "/api/files/upload",
            headers={"Origin": "https://evil.example"},
            files={"file": ("x.csv", b"a\n1")},
        ).status_code
        == 403
    )
    fid = upload(client)["file_id"]
    assert client.get(f"/api/files/{fid}/preview?page=0").status_code == 422
    assert client.get("/api/files/not-a-uuid/preview").status_code == 404
    record = store.get(fid, client.cookies.get("dataneat_session"))
    record["created"] = time.time() - 86401
    store.save(record)
    assert client.get(f"/api/files/{fid}/profile").status_code == 404


def test_resource_limits_pagination_and_cleanup(client, monkeypatch):
    import app.main as main

    monkeypatch.setattr(main, "MAX_UPLOAD", 5)
    assert (
        client.post(
            "/api/files/upload", files={"file": ("large.csv", b"a\n123456789")}
        ).status_code
        == 413
    )
    monkeypatch.setattr(main, "MAX_UPLOAD", 20 * 1024 * 1024)
    uploaded = upload(client, b"ID\n001\n002\n003\n")
    fid = uploaded["file_id"]
    paged = client.get(f"/api/files/{fid}/preview?page=2&page_size=1").json()
    assert paged["rows"] == [["002"]]
    path = store.path(fid)
    import os

    expired = time.time() - 86401
    os.utime(path, (expired, expired))
    store.purge()
    assert not path.exists()


def test_formula_headers_and_precision(client):
    result = upload(client, b'=header,Amount\n=1+1,"99999999999999999999.123456789"\n')
    fid = result["file_id"]
    ops = [
        {
            "operation": "normalize_numeric",
            "column": "Amount",
            "enabled": True,
            "parameters": {},
        }
    ]
    job = client.post(f"/api/files/{fid}/clean", json={"operations": ops}).json()[
        "job_id"
    ]
    wait_job(client, job)
    text = client.get(f"/api/jobs/{job}/download").text
    assert "'=header" in text and "'=1+1" in text
    assert "99999999999999999999.123456789" in text


def test_excel_formatted_identifier_and_encoding(client):
    workbook = Workbook()
    ws = workbook.active
    ws.append(["ID", "City"])
    ws.append([1, "Cairo"])
    ws["A2"].number_format = "000"
    buffer = io.BytesIO()
    workbook.save(buffer)
    result = client.post(
        "/api/files/upload", files={"file": ("book.xlsx", buffer.getvalue())}
    ).json()
    assert result["preview"]["rows"][0][0] == "001"
    assert result["profile"]["columns"][0]["detected_type"] == "identifier"
    result = client.post(
        "/api/files/upload",
        files={"file": ("unicode.csv", "ID,City\n001,القاهرة\n".encode("utf-16"))},
    ).json()
    assert result["preview"]["rows"][0][1] == "القاهرة"


def test_rate_limit_and_streamed_body(client):
    from collections import deque

    import app.main as main

    main.rate["testclient"] = deque([time.monotonic()] * 120)
    assert client.get("/api/health").status_code == 429
    main.rate.clear()
    import asyncio

    from app.services.limits import BodyLimitMiddleware
    from fastapi import HTTPException

    async def scenario():
        async def receive():
            return {"type": "http.request", "body": b"123456"}

        async def downstream(scope, receive, send):
            await receive()

        middleware = BodyLimitMiddleware(downstream, max_bytes=5)
        with pytest.raises(HTTPException) as exc:
            await middleware({"type": "http"}, receive, None)
        assert exc.value.status_code == 413

    asyncio.run(scenario())


def test_invalid_cleaning_parameters_are_rejected(client):
    fid = upload(client)["file_id"]
    invalid = [
        {
            "operation": "normalize_dates",
            "column": "Birth_Date",
            "parameters": {"interpretation": "guess"},
        },
        {
            "operation": "normalize_missing_values",
            "column": "City",
            "parameters": {"markers": "NA"},
        },
        {"operation": "remove_duplicates", "parameters": {"subset": "City"}},
        {"operation": "rename_column", "column": "City", "parameters": {"name": 123}},
    ]
    for operation in invalid:
        response = client.post(
            f"/api/files/{fid}/preview-cleaning",
            json={"operations": [{**operation, "enabled": True}]},
        )
        assert response.status_code == 422


def test_invalid_date_rows_are_paginated_isolated_and_correctable(client):
    result = upload(
        client,
        b"Event,Amount\n88-99-2089,$100\n31/02/2025,200 USD\n05/10/2025,300\n2025-10-05,EUR 400\n",
    )
    fid = result["file_id"]
    rows = client.get(
        f"/api/files/{fid}/invalid-dates",
        params={"column": "Event", "page_size": 1, "page": 2},
    ).json()
    assert rows["total"] == 2 and rows["rows"] == [
        {"row_index": 1, "spreadsheet_row": 3, "value": "31/02/2025"}
    ]
    outsider = TestClient(app)
    assert (
        outsider.get(f"/api/files/{fid}/invalid-dates?column=Event").status_code == 404
    )
    assert (
        client.get(f"/api/files/{fid}/invalid-dates?column=Amount").status_code == 422
    )
    dates = next(o for o in result["plan"] if o["operation"] == "resolve_invalid_dates")
    dates.update(
        enabled=True,
        parameters={"action": "manual", "replacements": {"0": "2025-10-06", "1": None}},
    )
    money = next(o for o in result["plan"] if o["operation"] == "normalize_currency")
    money.update(
        enabled=True,
        parameters={
            "currency": "USD",
            "format": "code",
            "assume_unlabelled": True,
            "symbol_aliases": ["$"],
        },
    )
    body = {"operations": [dates, money]}
    preview = client.post(f"/api/files/{fid}/preview-cleaning", json=body)
    assert preview.status_code == 200, preview.text
    assert preview.json()["after"]["rows"] == [
        ["2025-10-06", "USD 100"],
        [None, "USD 200"],
        ["05/10/2025", "USD 300"],
        ["2025-10-05", "EUR 400"],
    ]
    job = client.post(f"/api/files/{fid}/clean", json=body).json()["job_id"]
    wait_job(client, job)
    results = client.get(f"/api/jobs/{job}/results").json()
    assert results["validation"]["invalid_dates_resolved"] == 2
    assert (
        results["validation"]["unresolved_issues"] > 0
    )  # EUR and ambiguity still need review.
    assert client.get(f"/api/files/{fid}/preview").json()["rows"][0] == [
        "88-99-2089",
        "$100",
    ]
    for op in (
        {
            **dates,
            "parameters": {"action": "manual", "replacements": {"0": "88-99-2089"}},
        },
        {**money, "parameters": {"currency": "USD", "assume_unlabelled": "yes"}},
        {**money, "parameters": {"currency": "USD", "symbol_aliases": ["EUR"]}},
    ):
        assert (
            client.post(
                f"/api/files/{fid}/preview-cleaning", json={"operations": [op]}
            ).status_code
            == 422
        )


def test_percentage_api_formats_preserve_original_and_meaning(client):
    result = upload(client, b"Rate\n10%\n0.2\n30%\n")
    fid = result["file_id"]
    assert result["profile"]["columns"][0]["detected_type"] == "percentage"
    op = next(o for o in result["plan"] if o["operation"] == "normalize_percentage")
    op.update(
        enabled=True,
        parameters={"format": "with_sign", "bare_interpretation": "fraction"},
    )
    preview = client.post(
        f"/api/files/{fid}/preview-cleaning", json={"operations": [op]}
    ).json()
    assert preview["after"]["rows"] == [["10%"], ["20%"], ["30%"]]
    assert not any(
        i["issue"] == "percentage_inconsistency"
        for i in preview["validation"]["remaining_issues"]
    )
    op["parameters"]["format"] = "without_sign"
    job = client.post(f"/api/files/{fid}/clean", json={"operations": [op]}).json()[
        "job_id"
    ]
    wait_job(client, job)
    assert client.get(f"/api/jobs/{job}/results").json()["preview"]["rows"] == [
        ["10"],
        ["20"],
        ["30"],
    ]
    assert "10\n20\n30" in client.get(f"/api/jobs/{job}/download").text
    assert client.get(f"/api/files/{fid}/preview").json()["rows"] == [
        ["10%"],
        ["0.2"],
        ["30%"],
    ]
    op["parameters"]["bare_interpretation"] = "guess"
    assert (
        client.post(
            f"/api/files/{fid}/preview-cleaning", json={"operations": [op]}
        ).status_code
        == 422
    )


def test_complete_column_symbol_and_phone_choices(client):
    result = upload(
        client,
        b"Rate,Amount,Phone\n10%,$100,+201001234567\n20,200 EUR,201001234567\n0.2,300,01001234567\n",
    )
    fid = result["file_id"]
    assert any(
        i["issue"] == "phone_inconsistency"
        for c in result["profile"]["columns"]
        for i in c["issues"]
    )

    def op(name, column, parameters):
        return {
            "operation": name,
            "column": column,
            "enabled": True,
            "parameters": parameters,
        }

    actions = [
        op("normalize_percentage", "Rate", {"action": "add"}),
        op(
            "normalize_currency",
            "Amount",
            {"action": "add", "symbol": "$", "position": "suffix"},
        ),
        op("normalize_phone", "Phone", {"action": "international", "region": "EG"}),
    ]
    preview = client.post(
        f"/api/files/{fid}/preview-cleaning", json={"operations": actions}
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["after"]["rows"] == [
        ["10%", "100$", "+201001234567"],
        ["20%", "200$", "+201001234567"],
        ["0.2%", "300$", "+201001234567"],
    ]
    actions = [
        op("normalize_percentage", "Rate", {"action": "remove"}),
        op("normalize_currency", "Amount", {"action": "remove"}),
        op("normalize_phone", "Phone", {"action": "national", "region": "EG"}),
    ]
    job = client.post(f"/api/files/{fid}/clean", json={"operations": actions}).json()[
        "job_id"
    ]
    wait_job(client, job)
    assert client.get(f"/api/jobs/{job}/results").json()["preview"]["rows"] == [
        ["10", "100", "01001234567"],
        ["20", "200", "01001234567"],
        ["0.2", "300", "01001234567"],
    ]
    assert client.get(f"/api/files/{fid}/preview").json()["rows"][0] == [
        "10%",
        "$100",
        "+201001234567",
    ]
    unchanged = [
        op("normalize_percentage", "Rate", {"action": "preserve"}),
        op("normalize_currency", "Amount", {"action": "preserve"}),
        op("normalize_phone", "Phone", {"action": "preserve"}),
    ]
    assert (
        client.post(
            f"/api/files/{fid}/preview-cleaning", json={"operations": unchanged}
        ).json()["after"]["rows"]
        == result["preview"]["rows"]
    )
    for bad in [
        op("normalize_phone", "Phone", {"action": "international"}),
        op("normalize_phone", "Phone", {"action": "international", "region": "XX"}),
        op("normalize_currency", "Amount", {"action": "add", "symbol": "bad symbol"}),
        op("normalize_percentage", "Rate", {"action": "guess"}),
    ]:
        assert (
            client.post(
                f"/api/files/{fid}/preview-cleaning", json={"operations": [bad]}
            ).status_code
            == 422
        )


def test_imputation_preview_approval_execution_export_and_original(client):
    uploaded = upload(client, b"Amount,Category\n10,red\nNA,\n20,red\n")
    fid = uploaded["file_id"]
    suggestions = [
        o for o in uploaded["plan"] if o["operation"] == "impute_missing_values"
    ]
    assert len(suggestions) == 2 and all(not o["enabled"] for o in suggestions)
    assert uploaded["profile"]["columns"][0]["imputation"]["method"] == "mean"
    operations = [
        {
            "id": "amount",
            "column": "Amount",
            "operation": "impute_missing_values",
            "enabled": True,
            "parameters": {"method": "mean", "confirmed": False},
        },
        {
            "id": "category",
            "column": "Category",
            "operation": "impute_missing_values",
            "enabled": True,
            "parameters": {"method": "mode", "confirmed": False},
        },
    ]
    body = {"operations": operations}
    trial = client.post(f"/api/files/{fid}/preview-cleaning", json=body)
    assert trial.status_code == 200, trial.text
    assert trial.json()["after"]["rows"][1] == [15, "red"]
    assert trial.json()["after"]["is_imputed"] == [
        [False, False],
        [True, True],
        [False, False],
    ]
    assert trial.json()["validation"]["imputed_cells"] == 2
    blocked = client.post(f"/api/files/{fid}/clean", json=body)
    assert blocked.status_code == 422 and "Confirm imputation" in blocked.text
    for operation in operations:
        operation["parameters"]["confirmed"] = True
    response = client.post(f"/api/files/{fid}/clean", json=body)
    assert response.status_code == 202, response.text
    job = response.json()["job_id"]
    wait_job(client, job)
    result = client.get(f"/api/jobs/{job}/results?page=2&page_size=1").json()
    assert result["preview"]["is_imputed"] == [[True, True]]
    assert result["validation"]["missing_after"] == 0
    report = client.get(f"/api/jobs/{job}/download?format=report").json()
    assert report["operations"][0]["imputed_cells"][0]["spreadsheet_row"] == 3
    assert report["operations"][0]["distribution_before"]["numeric"]["mean"] == 15
    assert report["operations"][1]["method"] == "mode"
    csv = client.get(f"/api/jobs/{job}/download?format=csv").text
    assert "15.0,red" in csv
    workbook = load_workbook(
        io.BytesIO(client.get(f"/api/jobs/{job}/download?format=xlsx").content)
    )
    assert workbook.active["A3"].value == 15
    assert client.get(f"/api/files/{fid}/preview").json()["rows"][1] == ["NA", ""]
    outsider = TestClient(app)
    assert outsider.get(f"/api/jobs/{job}/results").status_code == 404
    assert client.delete(f"/api/files/{fid}").status_code == 200


def test_imputation_invalid_dependency_and_parameters_return_errors(client):
    uploaded = upload(
        client, b"Amount,Time\n10,05/10/2025\nNA,06/10/2025\n20,07/10/2025\n"
    )
    path = f"/api/files/{uploaded['file_id']}/preview-cleaning"
    operation = {
        "column": "Amount",
        "operation": "impute_missing_values",
        "enabled": True,
        "parameters": {
            "method": "linear_interpolation",
            "order_by": "Time",
            "single_entity": True,
        },
    }
    response = client.post(path, json={"operations": [operation]})
    assert (
        response.status_code == 422 and "Approve safe date conversion" in response.text
    )
    operation["parameters"] = {"method": "group_mean"}
    assert client.post(path, json={"operations": [operation]}).status_code == 422
    operation["parameters"] = {"method": "knn"}
    assert client.post(path, json={"operations": [operation]}).status_code == 422


def test_imputation_provenance_survives_rename_and_duplicate_removal(client):
    uploaded = upload(client, b"Amount,Group\n10,A\nNA,A\n10,A\n")
    operations = [
        {
            "column": "Amount",
            "operation": "impute_missing_values",
            "enabled": True,
            "parameters": {"method": "mean"},
        },
        {
            "column": "Amount",
            "operation": "rename_column",
            "enabled": True,
            "parameters": {"name": "Value"},
        },
        {"operation": "remove_duplicates", "enabled": True, "parameters": {}},
    ]
    data = client.post(
        f"/api/files/{uploaded['file_id']}/preview-cleaning",
        json={"operations": operations},
    ).json()
    assert data["after"]["columns"] == ["Value", "Group"]
    # String observations remain strings; a numeric estimate is distinct without an approved conversion.
    assert data["after"]["is_imputed"] == [[False, False], [True, False]]
    operations[0]["parameters"]["method"] = "mode"
    mode = client.post(
        f"/api/files/{uploaded['file_id']}/preview-cleaning",
        json={"operations": operations},
    ).json()
    # An identical copied mode is removed; the retained observed cell is not marked imputed.
    assert mode["after"]["is_imputed"] == [[False, False]]
