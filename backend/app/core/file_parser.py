import csv
import io
import re
import zipfile
from pathlib import Path

import pandas as pd
from charset_normalizer import from_bytes
from openpyxl import load_workbook

MAX_ROWS = 100000
MAX_COLUMNS = 200


def unique_headers(headers):
    seen, output = set(), []
    for index, value in enumerate(headers):
        base = str(value) if value not in (None, "") else f"Column_{index + 1}"
        name, suffix = base, 2
        while name in seen:
            name = f"{base}_{suffix}"
            suffix += 1
        seen.add(name)
        output.append(name)
    return output


def parse_file(data, filename, sheet=None):
    extension = Path(filename).suffix.lower()
    sheets = []
    if extension == ".csv":
        best = from_bytes(data).best()
        if not best:
            raise ValueError("Cannot detect CSV encoding")
        text = str(best)
        if "\x00" in text:
            raise ValueError("Invalid CSV contents")
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)
        try:
            headers = next(reader)
        except StopIteration:
            raise ValueError("The file is empty")
        if not headers or len(headers) > MAX_COLUMNS:
            raise ValueError("File must have 1–200 columns")
        rows = []
        for row in reader:
            if not row:
                continue
            if len(row) != len(headers):
                raise ValueError("CSV rows have inconsistent field counts")
            rows.append(row)
            if len(rows) > MAX_ROWS:
                raise ValueError("Maximum 100,000 rows supported")
        metadata = {"encoding": best.encoding, "delimiter": delimiter}
    elif extension == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("Workbook expands beyond the safety limit")
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            sheets = workbook.sheetnames
            if sheet and sheet not in sheets:
                raise ValueError("Unknown worksheet")
            selected = sheet or sheets[0]
            worksheet = workbook[selected]
            if worksheet.max_row > MAX_ROWS + 1 or worksheet.max_column > MAX_COLUMNS:
                raise ValueError("Workbook dimensions exceed limits")
            iterator = worksheet.iter_rows()
            headers = [cell.value for cell in next(iterator, [])]

            def cell_value(cell):
                value = cell.value
                if (
                    isinstance(value, int)
                    and not isinstance(value, bool)
                    and re.fullmatch(r"0{2,}", cell.number_format or "")
                ):
                    return str(value).zfill(len(cell.number_format))
                return value

            rows = []
            for row in iterator:
                rows.append([cell_value(cell) for cell in row])
                if len(rows) > MAX_ROWS:
                    raise ValueError("Maximum 100,000 rows supported")
            workbook.close()
            metadata = {"worksheet": selected}
        except (zipfile.BadZipFile, KeyError):
            raise ValueError("Invalid XLSX workbook")
    else:
        raise ValueError("Only CSV and XLSX files are supported")
    if not headers or len(headers) > MAX_COLUMNS:
        raise ValueError("File must have 1–200 columns")
    if not rows:
        raise ValueError("File has no data rows")
    if len(rows) > MAX_ROWS:
        raise ValueError("Maximum 100,000 rows supported")
    names = unique_headers(headers)
    return pd.DataFrame(rows, columns=names, dtype=object), {
        "sheets": sheets,
        "header_adjustments": [
            {"original": str(a), "name": b}
            for a, b in zip(headers, names)
            if str(a) != b
        ],
        **metadata,
    }
